"""
Chay: python tao_word.py
Yeu cau: pip install python-docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Cài font mặc định ──────────────────────────────────────────────────
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

# ── Hàm tiện ích ───────────────────────────────────────────────────────
def add_title(text, level=1):
    p = doc.add_heading(text, level=level)
    run = p.runs[0]
    run.font.name = "Times New Roman"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x0A, 0x29, 0x66)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xAA)
    elif level == 3:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_step(step_name, description):
    """In tên bước đậm + nội dung mô tả trên cùng dòng / dòng tiếp theo."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.3)
    run_title = p.add_run(f"{step_name}: ")
    run_title.bold = True
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(12)
    run_desc = p.add_run(description)
    run_desc.font.name = "Times New Roman"
    run_desc.font.size = Pt(12)

def add_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(f"⚠ Lưu ý: {text}")
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)

def add_separator():
    doc.add_paragraph("─" * 60)

# ══════════════════════════════════════════════════════════════════════
# TIÊU ĐỀ CHÍNH
# ══════════════════════════════════════════════════════════════════════
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run("CÁC BƯỚC KÝ SỐ VÀ XÁC THỰC\nTRONG HỆ THỐNG ANM (RSA-SHA256)")
run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x07, 0x1D, 0x4F)

doc.add_paragraph()  # dòng trống

# ══════════════════════════════════════════════════════════════════════
# PHẦN I — KÝ VĂN BẢN
# ══════════════════════════════════════════════════════════════════════
add_title("PHẦN I — KÝ VĂN BẢN (Text Signing)", level=1)

add_title("A. Quy trình ký", level=2)

add_step(
    "Bước 1 — Chuẩn hóa văn bản",
    "Văn bản đầu vào được trim, chuẩn hóa khoảng trắng và encode sang UTF-8 "
    "để đảm bảo xử lý nhất quán trước khi thực hiện các phép toán mật mã."
)

add_step(
    "Bước 2 — Tính hash SHA-256",
    "Áp dụng hàm băm SHA-256 lên văn bản đã chuẩn hóa (encode UTF-8), "
    "cho ra chuỗi hex 64 ký tự (256 bit) đại diện duy nhất cho nội dung. "
    "Thay đổi dù chỉ 1 ký tự cũng tạo ra hash hoàn toàn khác."
)

add_step(
    "Bước 3 — Ký RSA bằng Private Key",
    "Thư viện cryptography thực hiện: SHA256(text) → PKCS1v15 padding → "
    "mã hóa bằng private key d theo công thức: Signature = Padded^d mod N. "
    "Padding scheme: PKCS#1 v1.5. Hash function: SHA-256."
)

add_step(
    "Bước 4 — Encode Base64",
    "Chữ ký nhị phân (256 bytes) được encode sang Base64 để lưu trữ "
    "và truyền tải dưới dạng chuỗi ASCII an toàn."
)

add_step(
    "Bước 5 — Đóng gói thành file JSON",
    "Toàn bộ thông tin (nội dung gốc, hash, chữ ký, thuật toán, key_id, "
    "thời gian ký) được đóng gói vào file .json định dạng RSA_SIGNED_TEXT "
    "để người dùng tải về và dùng cho xác thực sau này."
)

add_title("B. Quy trình xác thực văn bản", level=2)

add_step(
    "Bước 1 — Đọc & kiểm tra cấu trúc gói JSON",
    "Tải lên file .json đã ký, kiểm tra các trường bắt buộc: format = "
    "RSA_SIGNED_TEXT, version = 1, hash_algorithm = SHA-256, "
    "signature_algorithm = RSA-SHA256, và các trường content/key_id/hash/signature."
)

add_step(
    "Bước 2 — Kiểm tra toàn vẹn nội dung (Hash check)",
    "Tính lại SHA-256 của trường content trong gói JSON, "
    "so sánh với giá trị hash đã lưu lúc ký. "
    "Nếu không khớp → nội dung đã bị sửa sau khi ký."
)

add_step(
    "Bước 3 — Xác thực chữ ký RSA",
    "Dùng public_key.verify() với: chữ ký Base64 đã decode, nội dung văn bản, "
    "PKCS1v15 padding, SHA-256. Bên trong: RSA decrypt bằng public key e, "
    "so sánh hash giải mã được với SHA256(text) hiện tại. "
    "Nếu không khớp → raise InvalidSignature."
)

add_step(
    "Bước 4 — Kết luận tổng hợp",
    "Kết quả cuối = signature_valid AND hash_valid. "
    "Nếu cả hai đều True → Chữ ký hợp lệ, nội dung chưa bị sửa. "
    "Nếu hash_valid = False → Nội dung bị sửa sau khi ký. "
    "Nếu signature_valid = False → Chữ ký sai hoặc sai khóa."
)

# ══════════════════════════════════════════════════════════════════════
# PHẦN II — KÝ PDF
# ══════════════════════════════════════════════════════════════════════
doc.add_page_break()

add_title("PHẦN II — KÝ FILE PDF (PDF Signing)", level=1)

add_title("A. Quy trình ký PDF", level=2)

add_step(
    "Bước 1 — Đọc file PDF dưới dạng bytes thô",
    "Toàn bộ nội dung file PDF được đọc vào bộ nhớ dưới dạng bytes nhị phân "
    "(không parse nội dung). Đây là dữ liệu gốc chưa qua bất kỳ chỉnh sửa nào."
)

add_step(
    "Bước 2 — Hash toàn bộ bytes PDF gốc bằng SHA-256",
    "SHA-256 được áp dụng lên toàn bộ binary của file PDF, "
    "cho ra giá trị hex 64 ký tự. Khác với ký văn bản: hash "
    "toàn bộ binary thay vì hash nội dung chữ."
)

add_step(
    "Bước 3 — Ký hash của PDF bằng RSA Private Key",
    "Chỉ ký 32 bytes hash (không ký cả file lớn). "
    "Thực hiện: RSA Sign(bytes.fromhex(pdf_hash), PKCS1v15, SHA-256). "
    "Lưu ý: do đầu vào là hex string của hash, nên thực chất là SHA256 hai lần (double-hash)."
)

add_step(
    "Bước 4 — Tạo tem hiển thị (Visual Stamp) lên trang PDF",
    "Tạo một layer PDF trong suốt chứa: nhãn DIGITALLY SIGNED, "
    "tên người ký, thời gian ký, và ảnh chữ ký (PNG/JPG) nếu có. "
    "Layer này được merge vào trang PDF được chỉ định. "
    "Tem chỉ có tác dụng hiển thị, KHÔNG tham gia vào toán học RSA."
)

add_step(
    "Bước 5 — Nhúng chữ ký số vào PDF Metadata",
    "Các trường metadata ẩn được ghi vào file PDF: "
    "/RSA_Signature (Base64), /RSA_Original_SHA256 (hex hash gốc), "
    "/RSA_Signature_Algorithm, /RSA_KeyID, /RSA_Signer, /RSA_SignedAt. "
    "Metadata này có thể xem qua Adobe Reader (File Properties)."
)

add_note(
    "Chữ ký số nằm trong metadata ẩn. Ảnh chữ ký hiển thị trên trang "
    "chỉ là hình ảnh trang trí, không có giá trị xác thực mật mã."
)

add_title("B. Quy trình xác thực PDF", level=2)

add_step(
    "Bước 1 — Đọc metadata ẩn từ file PDF tải lên",
    "Dùng PdfReader để đọc dictionary metadata của file PDF. "
    "Lấy hai trường quan trọng: /RSA_Signature và /RSA_Original_SHA256."
)

add_step(
    "Bước 2 — Kiểm tra sự tồn tại của metadata chữ ký",
    "Nếu không tìm thấy /RSA_Signature hoặc /RSA_Original_SHA256 "
    "→ kết luận ngay: file PDF không có chữ ký RSA của hệ thống."
)

add_step(
    "Bước 3 — Xác thực chữ ký RSA với hash gốc lưu trong metadata",
    "Thực hiện: public_key.verify(decode(signature), bytes.fromhex(original_hash), "
    "PKCS1v15, SHA256). Nếu thành công → chữ ký hợp lệ với hash đã lưu. "
    "Nếu raise InvalidSignature → metadata bị can thiệp hoặc sai khóa."
)

add_step(
    "Bước 4 — Trả về kết quả xác thực",
    "Trả về: valid (True/False), message (người ký & thời gian hoặc lý do lỗi), "
    "hash (giá trị hash gốc), verified_at (thời điểm xác thực)."
)

add_note(
    "Phiên bản hiện tại CHƯA tính lại hash từ nội dung PDF thực tế khi xác thực. "
    "Do đó chỉ phát hiện được metadata bị sửa, chưa phát hiện được "
    "nếu nội dung các trang PDF bị chỉnh sửa mà giữ nguyên metadata."
)

# ══════════════════════════════════════════════════════════════════════
# BẢNG SO SÁNH
# ══════════════════════════════════════════════════════════════════════
doc.add_page_break()
add_title("PHẦN III — SO SÁNH KÝ VĂN BẢN VÀ KÝ PDF", level=1)

table = doc.add_table(rows=8, cols=3)
table.style = "Table Grid"

headers = ["Tiêu chí", "Ký Văn bản", "Ký PDF"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    run = cell.paragraphs[0].add_run(h)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

rows_data = [
    ("Đầu vào ký",         "text.encode('utf-8')",          "bytes.fromhex(pdf_hash)"),
    ("Hash lưu ở đâu",     "File .json",                    "PDF Metadata (/RSA_Original_SHA256)"),
    ("Chữ ký lưu ở đâu",   "File .json",                    "PDF Metadata (/RSA_Signature)"),
    ("Visual stamp",        "Không có",                      "Tem hiển thị trên trang PDF"),
    ("Phát hiện sửa nội dung", "Có (hash_valid check)",     "Chưa đầy đủ (chỉ check metadata)"),
    ("File đầu ra",         "signed_text_*.json",            "signed_*.pdf"),
    ("Thư viện chính",      "cryptography + hashlib",        "cryptography + pypdf + reportlab"),
]

for i, (c1, c2, c3) in enumerate(rows_data, start=1):
    for j, text in enumerate([c1, c2, c3]):
        cell = table.rows[i].cells[j]
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

# ══════════════════════════════════════════════════════════════════════
# LƯU FILE
# ══════════════════════════════════════════════════════════════════════
output_path = "Chu_Ky_So_RSA_Cac_Buoc.docx"
doc.save(output_path)
print(f"✅ Đã tạo file: {output_path}")
